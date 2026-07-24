import os

def extract_text_from_file(file_path):
    """
    Extracts plain text content from uploaded PDF, DOCX, or TXT files.
    Includes fallbacks for corrupted or unreadable documents.
    """
    if not os.path.exists(file_path):
        return ""

    ext = file_path.rsplit('.', 1)[-1].lower() if '.' in file_path else ""

    if ext == 'pdf':
        return _extract_from_pdf(file_path)
    elif ext in ['docx', 'doc']:
        return _extract_from_docx(file_path)
    elif ext in ['txt', 'md']:
        return _extract_from_txt(file_path)

    return ""

def _extract_from_pdf(file_path):
    text = ""
    try:
        from pypdf import PdfReader
        reader = PdfReader(file_path)
        for page in reader.pages:
            t = page.extract_text()
            if t:
                text += t + "\n"
    except Exception as e:
        print(f"[FileParser] pypdf error: {e}. Trying fallback...")
        try:
            import PyPDF2
            reader = PyPDF2.PdfReader(file_path)
            for page in reader.pages:
                t = page.extract_text()
                if t:
                    text += t + "\n"
        except Exception as ex:
            print(f"[FileParser] PyPDF2 error: {ex}")
    
    return text.strip()

def _extract_from_docx(file_path):
    text = ""
    try:
        import docx
        doc = docx.Document(file_path)
        for para in doc.paragraphs:
            if para.text:
                text += para.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        text += cell.text + " "
                text += "\n"
    except Exception as e:
        print(f"[FileParser] python-docx error: {e}")
    
    return text.strip()

def _extract_from_txt(file_path):
    try:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read().strip()
    except Exception as e:
        print(f"[FileParser] TXT read error: {e}")
        return ""
